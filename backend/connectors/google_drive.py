"""Conector de carpeta de Google Drive.

La idea, en una línea: **Kino suelta el archivo del SYS en una carpeta de Drive
y el sistema lo toma solo**, sin que nadie tenga que entrar a la app a subirlo.

Es un conector más, igual que `local_file` o `google_sheets`: el motor no se
entera de dónde vino el archivo. La conexión guarda el id de la carpeta y en
cada corrida se busca el **archivo más reciente** que haya adentro.

Decisiones:
- **El más reciente, no "el que se llame X".** Así Kino puede subir
  `BASE-SYS 20-oct.xlsx` y mañana `BASE-SYS 21-oct.xlsx` sin configurar nada.
- **Se ignoran los que no son fuente**: los `.md` del proyecto que viven en la
  misma carpeta y los temporales que Excel deja mientras el archivo está
  abierto (`~$...`), que si no serían "el más reciente" y romperían la corrida.
- **El parseo se delega en `LocalFileConnector`**: ya sabe leer xlsx/xls/csv y
  arregla los SKU que Excel convierte en fecha (`9051-3`). Acá solo se baja el
  archivo; duplicar ese parser sería pedir que se desincronicen.
- Usa la **misma cuenta de servicio** que ya lee la Maestra. Para que funcione,
  la carpeta tiene que estar compartida con el correo de esa cuenta y la API de
  Drive habilitada en el proyecto de Google.
"""
import io
import json
import os
from typing import Any, Dict, List, Optional, Tuple

from .base import BaseConnector
from .local_file import LocalFileConnector

EXTENSIONES = (".csv", ".xlsx", ".xls")
SCOPES = ["https://www.googleapis.com/auth/drive.readonly"]


def es_archivo_fuente(nombre: str) -> bool:
    """¿Este archivo de la carpeta sirve como fuente de datos?

    Deja afuera los `.md` del proyecto (que viven en la misma carpeta) y los
    temporales de Excel abierto (`~$BASE-SYS.xlsx`), que son más nuevos que el
    archivo de verdad y se llevarían el puesto de "el más reciente"."""
    n = (nombre or "").strip()
    if not n or n.startswith("~$"):
        return False
    return n.lower().endswith(EXTENSIONES)


def elegir_archivo(archivos: List[Dict[str, Any]]) -> Optional[Dict[str, Any]]:
    """El más reciente de los que sirven como fuente, o None si no hay ninguno."""
    utiles = [a for a in archivos if es_archivo_fuente(a.get("name", ""))]
    if not utiles:
        return None
    return max(utiles, key=lambda a: a.get("modifiedTime") or "")


def _credenciales():
    """Mismas credenciales que usa la Maestra, pero con permiso de lectura de
    Drive. Si no están configuradas, devuelve None y el error se explica arriba."""
    from google.oauth2 import service_account

    crudo = os.getenv("GOOGLE_CREDENTIALS_JSON")
    if crudo:
        return service_account.Credentials.from_service_account_info(json.loads(crudo), scopes=SCOPES)
    archivo = os.getenv("GOOGLE_CREDENTIALS_FILE", "../credentials.json")
    if os.path.exists(archivo):
        return service_account.Credentials.from_service_account_file(archivo, scopes=SCOPES)
    return None


class GoogleDriveConnector(BaseConnector):
    """Lee el archivo más reciente de una carpeta de Drive."""

    def __init__(self, connection_config: Dict[str, Any]):
        super().__init__(connection_config)
        self.folder_id = (self.config.get("drive_folder_id") or "").strip()
        self._cache: Optional[Tuple[Dict[str, Any], bytes]] = None

    # ------------------------------------------------------------------ red
    def _servicio(self):
        from googleapiclient.discovery import build

        creds = _credenciales()
        if not creds:
            raise ValueError(
                "Faltan las credenciales de Google en el servidor "
                "(GOOGLE_CREDENTIALS_JSON). Son las mismas que usa la Maestra."
            )
        return build("drive", "v3", credentials=creds)

    def listar_archivos(self) -> List[Dict[str, Any]]:
        """Lo que hay en la carpeta. Método aparte para poder probar sin red."""
        if not self.folder_id:
            raise ValueError("Esta conexión no tiene configurada la carpeta de Drive.")
        resp = self._servicio().files().list(
            q=f"'{self.folder_id}' in parents and trashed = false",
            fields="files(id, name, modifiedTime, mimeType, size)",
            orderBy="modifiedTime desc",
            pageSize=100,
        ).execute()
        return resp.get("files", [])

    def descargar(self, file_id: str) -> bytes:
        """Bytes del archivo. Método aparte para poder probar sin red."""
        from googleapiclient.http import MediaIoBaseDownload

        buffer = io.BytesIO()
        pedido = self._servicio().files().get_media(fileId=file_id)
        bajada = MediaIoBaseDownload(buffer, pedido)
        terminado = False
        while not terminado:
            _, terminado = bajada.next_chunk()
        return buffer.getvalue()

    # --------------------------------------------------------------- lógica
    def archivo_actual(self) -> Tuple[Dict[str, Any], bytes]:
        """(info, contenido) del archivo más reciente. Se cachea por instancia:
        leer las hojas y después los datos no baja el archivo dos veces."""
        if self._cache is not None:
            return self._cache
        elegido = elegir_archivo(self.listar_archivos())
        if not elegido:
            raise ValueError(
                "La carpeta de Drive no tiene ningún archivo .xlsx, .xls o .csv. "
                "Subí el archivo del SYS y volvé a correr."
            )
        self._cache = (elegido, self.descargar(elegido["id"]))
        return self._cache

    def _lector(self) -> LocalFileConnector:
        """El parseo lo hace el mismo lector de archivos subidos a mano, para
        que los dos caminos se comporten igual (incluido el arreglo de los SKU
        que Excel guarda como fecha)."""
        info, contenido = self.archivo_actual()
        return LocalFileConnector({"file_name": info["name"], "file_content": contenido})

    # ------------------------------------------------------------ interfaz
    def fetch_data(self, source_path: str) -> List[Dict[str, Any]]:
        return self._lector().fetch_data(source_path)

    def normalize_data(self, raw_data: List[Dict[str, Any]], field_mappings: Dict[str, str]) -> List[Dict[str, Any]]:
        return self._lector().normalize_data(raw_data, field_mappings)

    def metadata(self) -> Dict[str, List[str]]:
        """{hoja: [columnas]} del archivo más reciente, para el mapeo en la UI."""
        import openpyxl
        import pandas as pd

        info, contenido = self.archivo_actual()
        nombre = info["name"].lower()
        if nombre.endswith(".csv"):
            df = pd.read_csv(io.BytesIO(contenido), nrows=0)
            return {"CSV Data": df.columns.tolist()}
        wb = openpyxl.load_workbook(io.BytesIO(contenido), read_only=True)
        salida = {}
        for hoja in wb.sheetnames:
            ws = wb[hoja]
            try:
                encabezados = next(ws.iter_rows(max_row=1, values_only=True))
            except StopIteration:
                encabezados = ()
            salida[hoja] = [("" if v is None else str(v)) for v in encabezados]
        return salida

    def test_connection(self) -> Tuple[bool, str]:
        try:
            archivos = self.listar_archivos()
        except Exception as e:
            return False, (
                f"No se pudo leer la carpeta: {str(e)[:200]}. Revisá que la carpeta esté "
                f"compartida con el correo de la cuenta de servicio de Google."
            )
        elegido = elegir_archivo(archivos)
        if not elegido:
            return False, (f"La carpeta se lee bien ({len(archivos)} archivo(s)), pero no hay "
                           f"ninguno .xlsx, .xls o .csv para usar como fuente.")
        return True, f"Carpeta OK. El más reciente es '{elegido['name']}'."
